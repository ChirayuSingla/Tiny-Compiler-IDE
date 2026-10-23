"""Build the editable Review 2 report from recorded tests and genuine screenshots.

Run tools/verify_project.py and tools/verify_gui.py before rebuilding evidence.
This script uses python-docx (requirements-dev.txt) and does not invent results.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report/Tiny_Programming_Language_Compiler_Report.docx"


def main():
    evidence = json.loads((ROOT / "docs/test-results.json").read_text(encoding="utf-8"))
    native = json.loads((ROOT / "screenshots/verification.json").read_text(encoding="utf-8"))
    library = json.loads((ROOT / "examples/library.json").read_text(encoding="utf-8"))
    if not evidence["success"] or evidence["skipped"]:
        raise SystemExit("Refresh verification evidence: full suite must pass without skips.")
    d = Document()
    section = d.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.72)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = d.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = d.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.06
    d.styles["Title"].font.size = Pt(26)
    d.styles["Heading 1"].font.size = Pt(17)
    d.styles["Heading 2"].font.size = Pt(12)
    for name in ("Heading 1", "Heading 2"):
        d.styles[name].paragraph_format.space_before = Pt(8)
        d.styles[name].paragraph_format.space_after = Pt(6)
    d.core_properties.title = "Tiny Programming Language Compiler IDE"
    d.core_properties.subject = "Review 2 implementation and verification report"
    d.core_properties.author = "Tiny Compiler Project"
    # The stock Word template can put a blue border under Title. Remove its
    # paragraph-border definitions rather than relying on font formatting.
    for border in list(d.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)

    def p(text, style=None):
        return d.add_paragraph(text, style)

    def h(text):
        d.add_heading(text, level=1)

    def sub(text):
        d.add_heading(text, level=2)

    def page(title):
        heading = d.add_heading(title, level=1)
        heading.paragraph_format.page_break_before = True

    def code(text):
        paragraph = p("")
        paragraph.paragraph_format.space_after = Pt(7)
        run = paragraph.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        return paragraph

    def table(headers, rows, widths=None):
        t = d.add_table(rows=1, cols=len(headers))
        t.autofit = False
        if widths:
            for col, width in zip(t.columns, widths):
                col.width = Inches(width)
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = text
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "233B53")
            cell._tc.get_or_add_tcPr().append(shade)
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
                if i % 2 == 1:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "F1F4F7")
                    cell._tc.get_or_add_tcPr().append(shade)
        for row_index, row in enumerate(t.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for col_index, cell in enumerate(row.cells):
                if widths:
                    cell.width = Inches(widths[col_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    item = OxmlElement("w:" + edge)
                    item.set(qn("w:val"), "single")
                    item.set(qn("w:sz"), "4")
                    item.set(qn("w:color"), "D9D9D9")
                    borders.append(item)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    item = OxmlElement("w:" + edge)
                    item.set(qn("w:w"), "75")
                    item.set(qn("w:type"), "dxa")
                    margins.append(item)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        p("").paragraph_format.space_after = Pt(0)
        return t

    def picture(name, caption, width=6.8):
        path = ROOT / "screenshots" / name
        if not path.exists():
            raise FileNotFoundError("Genuine screenshot required: " + str(path))
        paragraph = p("")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        paragraph = p(caption)
        paragraph.paragraph_format.space_after = Pt(8)
        for run in paragraph.runs:
            run.italic = True
            run.font.size = Pt(9)

    # Page 1
    p("Tiny Programming Language Compiler IDE", "Title")
    p("Review 2 implementation report", "Subtitle")
    p("Verified on 19 October 2026", "Subtitle")
    sub("Project purpose and outcome")
    p("The project is a working desktop environment for writing Tiny programs and inspecting "
      "their compilation. It reuses the existing Python compiler and virtual machine. A source "
      "editor, eight selectable stage views and a persistent terminal let the user follow one "
      "program from characters to execution without rerunning it to change views.")
    p("The completed application passed %d automated tests with no skips, including the 198 "
      "original compiler and CLI tests. A native Windows window was launched, interactive input "
      "was exercised and genuine widget captures were recorded. The application also provides "
      "a smart editor, program library, real compiler visualizations and conservative complexity "
      "estimates. This report records implementation, verification, limits and Review 2 mapping." % evidence["tests"])
    sub("Project details")
    table(["Field", "Details to complete before academic submission"], [
        ["Student names", "[Student names]"], ["Register numbers", "[Register numbers]"],
        ["Programme and branch", "[Programme and branch]"],
        ["Project guide", "[Guide name and designation]"],
        ["Institution", "[Institution name]"], ["Review date", "[Review date]"],
    ], [1.8, 5.0])
    sub("Objectives and scope")
    p("Provide an editable source workspace, real results for every compiler stage, responsive "
      "execution and input, useful diagnostics, and reproducible evidence. The target is the "
      "included stack virtual machine. Native machine-code generation, optimization, multi-file "
      "builds and debugger stepping are outside this implementation.")
    p("The repository defines its own typed Tiny dialect. It supports integers, booleans, "
      "block scopes, conditionals, loops, input and output. Its exact grammar and semantics are "
      "documented in docs/GRAMMAR.md and docs/LANGUAGE_SPEC.md.")

    # Page 2
    page("Architecture and compiler reuse")
    p("The GUI is a front end to the existing modules. Compiler passes have not been replaced "
      "with sample data or a second demonstration compiler. The CLI remains a separate entry "
      "point and continues to use the same backend.")
    code("Main thread: editor -> captured source -> CompilerWorker QThread\n"
         "Worker: lexer -> parser -> semantics and symbols -> IR -> VM code\n"
         "Worker: VM(result.code) -> queued input and buffered output\n"
         "Main thread: copied snapshots -> stage inspectors and statuses\n"
         "Main thread: output signals -> persistent terminal")
    table(["Layer", "Actual modules and responsibility"], [
        ["Lexical and syntax", "lexer.py produces positioned tokens; parser.py builds ast_nodes.py dataclasses."],
        ["Semantics", "semantic.py annotates types and resolves scoped symbols from symbol_table.py."],
        ["Lowering and target", "ir.py produces three-address instructions; codegen.py generates stack instructions."],
        ["Execution", "vm.py executes generated code; runtime.py supplies arithmetic and typed I/O."],
        ["Pipeline", "pipeline.py retains partial results and emits optional progress callbacks."],
        ["Desktop", "ide/window.py, editor.py and views.py implement widgets and source actions."],
        ["Worker and state", "ide/worker.py owns a run; session.py supplies snapshots, queue input and output buffering."],
    ], [1.5, 5.3])
    sub("Thread ownership and consistency")
    p("Edits trigger a 350 ms debounced real compilation with source, filename and revision ID. "
      "Run reuses successful current bytecode or compiles the captured source first, then executes "
      "that exact result.code. Deep copies cross Qt signals "
      "because semantic analysis annotates AST objects in place. Only the main thread edits "
      "widgets. Stage navigation changes a stacked-view index and does not invoke compilation.")
    sub("Result lifecycle")
    p("Results stay available after execution and after clearing the terminal. Editing marks "
      "preceding results stale until background compilation updates the inspectors. An earlier "
      "runtime retains its own source version, bytecode and input. Compilation and execution IDs "
      "prevent late results from crossing revisions. Concurrent VMs are blocked; editing and "
      "navigation stay available. Pages 15 and 16 document the real-time workflow.")
    sub("Tools and environment")
    p("Python %s, PySide6 6.12.0, standard-library unittest, and Windows were used for application "
      "verification. The report is editable DOCX generated with python-docx. Qt is not imported "
      "by the compiler API or CLI." % evidence["python"])

    # Page 3
    page("Desktop workspace and controls")
    p("Source is edited on the left, the selected inspector occupies the center, stage "
      "navigation is on the right, and the terminal remains below all panels. Splitters let "
      "the user resize the workspace. The screenshot below is an actual native Qt capture.")
    picture("02_ast.png", "Figure 1  Expandable AST from loop.tiny with the existing output still visible")
    table(["Action", "Shortcut", "Result"], [
        ["New Open Save", "Ctrl N / O / S", "UTF-8 source files and unsaved-change protection"],
        ["Save As", "Ctrl Shift S", "Save a separate source file from the File menu"],
        ["Compile", "F5", "Build stage results without running"],
        ["Run", "Ctrl Enter", "Reuse current compilation or compile first; execute VM code"],
        ["Stop", "Shift F5", "Cancel VM execution, input wait or between compiler passes"],
        ["Clear Terminal", "Ctrl L", "Clear terminal display while retaining stage results"],
        ["Program Library", "Toolbar or sidebar", "Load, run or inspect categorized examples"],
        ["Demo Mode", "Toolbar toggle", "Open the live pipeline when compiling or running"],
    ], [1.25, 1.5, 4.05])
    p("The editor paints line numbers, highlights Tiny keywords, numbers, strings and comments, "
      "and supports tab indentation. Samples open as unsaved copies. Token, symbol and AST rows "
      "can navigate to current source locations by double-click. Stale snapshots disable that "
      "location navigation to avoid pointing into a different program.")

    # Page 4
    page("Selectable compiler stages")
    table(["Stage", "Displayed result"], [
        ["Source Code", "Exact source snapshot for the selected run"],
        ["Lexical Analysis", "Token kind, lexeme, decoded value and line/column, including EOF"],
        ["Syntax Analysis", "Expandable AST with recovered partial tree when parsing fails"],
        ["Semantic Analysis", "Analyzer diagnostics, check summary and actual annotated AST"],
        ["Symbol Table", "Name, type, scope depth, storage identity, declaration and usage"],
        ["Intermediate Code", "Formatted three-address instructions from the annotated AST"],
        ["Target VM Code", "Indexed stack-machine instructions generated from that IR"],
        ["Execution Diagnostics", "Exact output, diagnostics, executed steps and final VM memory"],
    ], [1.65, 5.15])
    picture("01_tokens.png", "Figure 2  Token table generated by the original lexer for the retained compilation")
    p("Selecting a stage never dumps its contents into the terminal. AST types and resolved "
      "symbols reflect semantic annotations from the same run. This is an abstract syntax tree, "
      "not a concrete parse tree containing every punctuation token. Semantic analysis produces "
      "the symbol table; the sidebar exposes it as a separate inspection view.")
    p("Five additional sidebar views provide the live Pipeline, CFG, Execution Trace, Program "
      "Library and Complexity Analysis. They retain the original compiler-stage ordering and "
      "do not rerun code when selected. Pages 10 to 14 describe these extensions.")

    # Page 5
    page("Interactive input and execution")
    p("The VM runs in a worker thread. Its real InputReader validation is reused through a "
      "queue adapter. Input waits do not block the event loop, and stages remain selectable.")
    picture("03_input_wait.png", "Figure 3  GCD waits for terminal input while its generated VM instructions remain inspectable")
    picture("04_execution.png", "Figure 4  Supplying 84 36 executes the retained VM code and produces the expected GCD result")
    p("Whitespace-separated values can supply multiple reads. Invalid values produce runtime "
      "diagnostics. Output is buffered for responsive updates and flushed before prompts, "
      "errors and completion. Stop also interrupts an input wait; no global stdin replacement "
      "or forced thread termination is used.")

    # Page 6
    page("Errors progress and stale results")
    picture("05_semantic_error.png", "Figure 5  Type mismatch marks semantic analysis as error and retains earlier results and partial symbols")
    picture("06_stale_results.png", "Figure 6  Editing source visibly marks the retained compilation stale")
    p("Statuses distinguish pending, active, completed, waiting, partial, error, blocked, "
      "cancelled and not run. Errors stop unsafe downstream passes. Runtime failures preserve "
      "compiled code and preceding output. Fast compiler stages may finish before their active "
      "state is visually observable. Diagnostics include source positions and caret context "
      "where supplied by the backend.")

    # Page 7
    page("Executed tests and acceptance evidence")
    p("The original suite passed 198 tests before changes. The final suite passed %d tests "
      "with no skips in %.3f seconds. Recorded run: %s. The log is tests/latest-results.txt "
      "and machine-readable evidence is docs/test-results.json." % (
          evidence["tests"], evidence["seconds"], evidence["timestamp"]))
    table(["Test module", "Passed", "Coverage"], [
        [name + ".py", count, {
            "test_codegen_vm": "Stack code, runtime and differential arithmetic",
            "test_end_to_end": "CLI examples, flags and exit codes",
            "test_ide": "Stage API, widgets, worker execution and terminal input",
            "test_ir": "Three-address code and reference interpreter",
            "test_lexer": "Tokens, positions, comments and invalid characters",
            "test_parser": "Grammar, precedence and recovery",
            "test_semantics": "Types, scopes, declarations and warnings",
            "test_complexity": "Supported/ambiguous loop patterns and scalar-space model",
            "test_editor": "Grammar-aware completion and smart editor interaction",
            "test_program_library": "All library outputs, intentional errors and GUI actions",
            "test_visualizations": "Real CFG/VM trace, search, replay and stale estimates",
            "test_realtime": "Debounced stages, cache, revision guards and in-app input",
        }[name]] for name, count in evidence["passed_by_module"].items()
    ], [1.85, 0.6, 4.35])
    table(["Acceptance requirement", "Executed evidence"], [
        ["Same compilation", "Inspector values match retained objects; IR/code regeneration and independent VM output agree."],
        ["Success and errors", "All nine valid examples run through GUI worker; lexical, syntax, semantic and runtime failures exercised."],
        ["Input and responsiveness", "Int/bool input, multiple and blank lines, queued waits and output before prompts."],
        ["Stage navigation", "All selectors used during a paused compiler pass, actual VM input wait and after completion."],
        ["Persistent terminal", "Navigation preserves transcript; clear removes display but retains execution result."],
        ["Stale results and lifecycle", "Edits during execution, revert detection, stop/restart, close during input, Save/Open and samples."],
    ], [1.7, 5.1])
    sub("Native launch verification")
    p("Native Qt platform %s assertions covered sidebar navigation, input, errors, the library, "
      "graphs, trace, completion and complexity. They also verified automatic stages without execution "
      "and cached Run with sequential terminal input. The %d genuine captures and "
      "screenshots/verification.json record that run." % (native["platform"], len(native["screenshots"])))
    sub("What was not verified")
    p("Linux and macOS, human usability studies, full accessibility audits and large-program "
      "performance benchmarks were not run. File-dialog automation supplies paths rather than "
      "driving operating-system dialog chrome. These limits do not replace the executed widget "
      "and native-window tests.")

    # Page 8
    page("Review 2 mapping and presentation")
    table(["Review 2 item", "Evidence in this report and project"], [
        ["1 Project details", "Page 1 metadata placeholders remain editable for institutional details."],
        ["2 Introduction and need", "Page 1 explains interactive exploration of the compiler pipeline."],
        ["2 Objectives and solution", "Pages 1 to 3 describe scope, architecture and workspace."],
        ["3 Actual implementation", "Pages 2 to 6 and tiny_compiler/ide contain the working GUI and backend integration."],
        ["3 Tools and architecture", "Page 2 and docs/ARCHITECTURE.md describe actual modules and thread ownership."],
        ["3 Screenshots and results", "Figures are genuine application/widget captures; page 7 records executed tests."],
        ["3 Remaining work", "Page 9 lists current limitations and focused future work."],
        ["4 Novelty and differentiation", "Selectable retained stages, persistent I/O and snapshot consistency support teaching."],
        ["6 Presentation and questions", "Outline and viva prompts below. The prior Review 2 numbering skips item 5."],
        ["Advanced extensions", "Pages 10 to 14 cover visualizations, smart editor, library, Demo Mode and complexity."],
    ], [1.75, 5.05])
    sub("Contribution and differentiation")
    p("The contribution is the interactive inspection workflow over the existing compiler: "
      "retained stage results, independent terminal I/O, progress/error states and source "
      "staleness. It does not claim a new compiler algorithm. Compared with the previous "
      "CLI-only project, students can move among representations without rerunning the program "
      "or mixing every compiler listing with program output.")
    sub("Five minute demonstration")
    p("Introduce the learning problem and architecture. Run loop.tiny and compare tokens, AST, "
      "symbols, IR and stack instructions. Run gcd.tiny with 84 36 while navigating stages. "
      "Show semantic failure and stale results. Close with executed evidence and limitations.")
    sub("Viva prompts")
    p("Why an AST? It retains grammatical meaning without punctuation. Why IR? It separates "
      "language analysis from target generation. Why unique storage names? They distinguish "
      "shadowed variables. Why copy snapshots? Semantic analysis mutates AST annotations. "
      "Why a queue? The VM can await input without blocking the GUI event loop. How is one run "
      "verified? Inspectors use one retained CompilationResult and the VM executes its code.")

    # Page 9
    page("Setup limits and maintenance")
    sub("Windows installation and launch")
    p("Use 64-bit Python 3.10 or newer. Open PowerShell in the project folder. Activation is "
      "optional because the commands call the environment interpreter directly.")
    code("python -m venv .venv\n"
         ".\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n"
         ".\\.venv\\Scripts\\python.exe gui.py\n"
         ".\\.venv\\Scripts\\python.exe -m tiny_compiler.ide examples\\gcd.tiny\n"
         ".\\.venv\\Scripts\\python.exe tools\\verify_project.py\n"
         ".\\.venv\\Scripts\\python.exe tools\\verify_gui.py")
    sub("Source distribution")
    code("gui.py                         Desktop launcher\n"
         "main.py                        Existing CLI launcher\n"
         "tiny_compiler/                 Existing compiler and VM\n"
         "tiny_compiler/ide/             Window, editor, views, worker, session\n"
         "examples/                      Original and library source fixtures\n"
         "tests/                         Compiler, CLI and Qt tests\n"
         "docs/                          Architecture, language and test evidence\n"
         "screenshots/                   Native Qt captures and verification JSON\n"
         "report/                        Editable Word report\n"
         "tools/                         Verification, report and ZIP builders\n"
         "dist/Tiny-Compiler-IDE.zip      Source package without dependencies")
    sub("Known limitations and future work")
    p("The stack VM is the only generated target. There are no functions, arrays, floats, "
      "optimizer or live debugger stepping. Trace controls replay recorded history. The newest compilation and most recent execution are retained separately, "
      "and terminal/session state is not restored after closing. Compiler cancellation is "
      "between passes. Large sources or unbounded output may consume memory. Smaller windows "
      "can require scrolling or panel resizing. Other desktop platforms remain untested.")
    p("Focused future work could add session export, bounded terminal history, scalable table "
      "models and instruction stepping. The source ZIP requires Python and dependency "
      "installation on the destination machine; it is not a standalone executable.")
    sub("Maintenance and references")
    p("README.md contains launch commands and user workflow. docs/LANGUAGE_SPEC.md and "
      "docs/GRAMMAR.md define the implemented dialect. docs/ARCHITECTURE.md explains ownership "
      "and callbacks. docs/TESTING.md, test-results.json and latest-results.txt provide executed "
      "evidence. These project files, rather than external performance claims, are the technical "
      "references for this implementation report.")
    p("To rebuild this editable report, install requirements-dev.txt and run tools/build_report.py "
      "after refreshing verification evidence. Academic names, register numbers, institution "
      "and guide details on page 1 must be completed before submission.")

    # Pages 10 to 14: implemented extensions, with genuine native evidence.
    page("Compiler pipeline and control flow")
    p("The pipeline diagram colors actual stage statuses for the captured source. Clicking a "
      "node opens its inspector. CFG blocks and labelled edges come from the same generated IR, "
      "including conditional branches, short-circuit logic and loop back edges.")
    picture("08_live_pipeline.png", "Figure 7  Live pipeline while the real nested-loop VM waits for input")
    picture("09_control_flow.png", "Figure 8  Control-flow graph derived from the retained nested-loop intermediate code")
    p("The graph is structural, not a measured execution path. Both branch edges may be present "
      "even if a particular input takes one path. Nodes support selection, full instruction "
      "details, panning and zoom. No external graph-rendering dependency or hard-coded example "
      "graph is used. The AST remains expandable and symbols can be searched across columns.")

    page("Execution trace and demonstration mode")
    picture("10_execution_trace.png", "Figure 9  Real VM history with instruction pointers, post-state and source-line selection")
    p("An optional VM observer records completed or failed instructions, IP and next IP, source "
      "line, variable/temporary values, stack and output delta. These records are generated by "
      "the VM executing the retained bytecode. The existing CLI text trace remains compatible.")
    table(["Presentation control", "Actual behavior"], [
        ["Replay Step", "Select the next captured instruction and its state"],
        ["Replay Continue and Pause", "Advance through captured rows using a GUI timer, or pause replay"],
        ["Reset Replay", "Return to the first captured instruction without executing the program again"],
        ["Demo Mode", "Select the live pipeline on Compile/Run; all completed views stay available"],
    ], [2.0, 4.8])
    p("These controls replay history; they are not live debugger stepping or breakpoints. "
      "Replay highlights the corresponding source line only while the source is current, and "
      "pauses when its view is hidden. The terminal remains unchanged during replay.")
    sub("Capture limits")
    p("The GUI retains the first 2,000 instructions. Each row captures at most 256 memory "
      "values, 256 stack values and 4,096 output characters. Truncation is labelled and the "
      "program continues normally. The trace is intentionally bounded for educational use; "
      "retained output in the terminal and final execution result are separate.")

    page("Smart source editor")
    picture("12_smart_editor.png", "Figure 10  Editing changes the source while retained results are visibly marked stale")
    picture("13_completion_popup.png", "Figure 11  Genuine separate capture of the completion popup for the prefix wh", width=3.0)
    p("Completion follows supported Tiny grammar and actual analyzer symbols, filtered by "
      "lexical scope, declaration position and current declaration identity. It offers keywords, "
      "operators, read/write statements and editable templates. Tiny has no callable built-in "
      "functions; unsupported constructs such as for, arrays or print are not suggested.")
    p("Suggestions appear while typing or with Ctrl Space. Up/Down navigates, Enter/Tab accepts "
      "and Esc dismisses. Acceptance replaces only the typed prefix and never executes source. "
      "Parentheses, braces and double quotes are paired, blocks indent automatically, matching "
      "brackets are highlighted and standard undo/redo/clipboard shortcuts remain available.")
    p("A 350 ms debounced full compilation updates actual stage views, underlines diagnostics "
      "and supplies hover information. It does not execute source or write errors into the "
      "terminal while typing. Partial syntax uses separately copied recovered declarations "
      "for conservative completion. Revision guards reject late results. The service is not "
      "a complete language server.")

    page("Program library and educational examples")
    picture("07_program_library.png", "Figure 12  Dedicated categorized library with metadata, source preview and explicit actions")
    p("The library contains %d programs: %d valid examples and %d explicitly labelled error "
      "demonstrations. Every entry supplies difficulty, description, compiler concepts, source, "
      "sample input, expected output and intended diagnostic where appropriate." % (
          len(library), sum(not e["expected_error"] for e in library), sum(bool(e["expected_error"]) for e in library)))
    table(["Level", "Programs", "Coverage"], [
        ["Beginner", sum(e["difficulty"] == "Beginner" for e in library),
         "Greetings, variables, arithmetic, precedence, input, comparisons, conditions, sign and parity"],
        ["Intermediate", sum(e["difficulty"] == "Intermediate" for e in library),
         "Counters, sum, factorial, multiplication table, Fibonacci, primes, GCD and reversal"],
        ["Advanced", sum(e["difficulty"] == "Advanced" for e in library),
         "Nested loops/conditions, pattern output, compound calculations, combined input/control and errors"],
    ], [1.1, 0.75, 4.95])
    p("Load creates an unsaved editable copy. Run Example uses sample input through the actual "
      "terminal queue. Inspect Compilation builds without executing. Search covers names and "
      "concepts. Loading, editing and saving a personal copy preserve the original fixtures. "
      "All valid library outputs are verified through VM, TAC and GUI execution, and lexical, "
      "syntax, semantic and runtime demonstrations are checked for their intended failure.")
    p("A presentation can run Sum of first N numbers with input 5, inspect its O(n) estimate, "
      "then compare Nested loops and Number reversal. GCD demonstrates a correct executable "
      "program whose complexity is intentionally unresolved by this conservative analyzer.")

    page("Educational complexity analysis")
    picture("11_complexity.png", "Figure 13  Actual nested-loop analysis reports O(n squared) under an explicit scalar model")
    table(["Supported pattern", "Upper bound in runtime input magnitude"], [
        ["Monotone constant-step counter", "O(n), with constant start/threshold and invariant input bound"],
        ["Positive multiplication or valid division by factor at least two", "O(log n)"],
        ["Independent nested loops with reset counters", "Products such as O(n squared), O(m times n) and O(n log n)"],
        ["Literal-only fixed loop bounds", "O(1) relative to runtime input"],
        ["Scalar auxiliary storage", "O(1) scalar slots for a fixed program; actual symbol/temporary counts shown"],
    ], [3.0, 3.8])
    p("The analyzer uses the semantically annotated AST and symbol identities, never measured "
      "sample runtimes. It explains counter recurrence, nesting, bounds, assumptions and "
      "confidence. Ambiguous termination, compound conditions, changing bounds, conditional "
      "updates, dependent nesting and unsupported recurrences report Unable to determine.")
    p("The model counts scalar operations, not Python integer bit costs. O(1) scalar slots "
      "does not mean O(1) bits or bytes; integers may grow. Input/output buffers, IDE trace and "
      "compiler artifacts are excluded. Exact time and memory are not claimed. Measured VM "
      "instruction counts remain in Execution Results. The full supported subset and tests "
      "are documented in docs/COMPLEXITY.md and tests/test_complexity.py.")

    page("Real-time compilation and source versions")
    p("Typing starts a 350 ms debounce. A background QThread calls the same compiler through "
      "code generation, publishing real tokens, AST, semantic results, symbols, IR and VM "
      "instructions. CFG and complexity are derived from that result. Errors preserve earlier "
      "stages and block unsafe downstream work. No program is executed automatically.")
    picture("14_realtime_compilation.png", "Figure 14  Typing alone generated revision 8 VM code; the retained runtime still identifies revision 6")
    sub("Compile and Run behavior")
    p("Compile builds immediately. Run checks exact source and revision, reuses successful "
      "current bytecode when available and compiles first otherwise. Invalid current source "
      "cannot execute older code. Reusing a compilation creates a new execution ID. Source "
      "revision and worker-identity checks reject obsolete signals, including rapid edits.")
    sub("Editing during execution")
    p("The newest compilation and active/most recent execution are separate sessions. While "
      "an earlier VM waits for input, the editor can compile a newer revision. Stage inspectors "
      "follow the editor; runtime output, trace and pending input keep their original identity. "
      "Runtime views visibly identify earlier source. Terminal text and completed results "
      "survive sidebar navigation, source edits and completion.")
    sub("Executed acceptance evidence")
    p("Fifteen tests in test_realtime.py exercise terminal entry, sequential reads, output "
      "before input, Stop/restart, EOF/invalid input, all-stage debounce, cache reuse, stale-code "
      "rejection, rapid edits, navigation, responsive Qt timers, simultaneous source/runtime "
      "versions, clear/repeat-Run behavior and distinct transcript formats. All passed in the "
      "recorded full suite. docs/TESTING.md maps the eleven requested scenarios to exact tests.")

    page("Integrated terminal input demonstration")
    picture("15_terminal_input.png", "Figure 15  Cached current compilation runs in the actual VM and waits for terminal input")
    picture("16_terminal_results.png", "Figure 16  Entering 20 then 22 inside the IDE produces the genuine output sum = 42")
    p("The native verification typed both values into the integrated field with Enter. The "
      "VM pauses via a queue adapter around its existing typed InputReader. Requests include "
      "execution ID and sequence; a lock rejects duplicate, cancelled or outdated submissions. "
      "No external console, input argument or alternate interpreter is required.")
    p("The terminal streams output, prompts, input echo, errors and lifecycle records in distinct "
      "colors. Separate visible labels show compilation and runtime status. End Input sends "
      "EOF and reports input exhaustion. Stop cancels a waiting read safely. Clear Terminal "
      "retains compiler/runtime state and a pending read. Automatic compiler diagnostics stay "
      "in editor/inspectors; explicit Compile/Run errors also appear in the terminal.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Tiny Compiler IDE  |  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    REPORT.parent.mkdir(exist_ok=True)
    d.save(REPORT)
    print("Wrote", REPORT)
    print("Evidence: %d tests, zero skips; native platform %s" % (evidence["tests"], native["platform"]))


if __name__ == "__main__":
    main()
